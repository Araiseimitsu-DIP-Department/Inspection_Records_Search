"""Read-only PostgreSQL smoke check with Access imports blocked and no GUI."""
from __future__ import annotations

import argparse
import datetime as dt
import importlib.abc
from pathlib import Path
import sys
import tempfile
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))


class BlockAccess(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split(".")[0] in {"pyodbc", "pypyodbc", "win32com"} or fullname.endswith("access_gateway"):
            raise ImportError("Access imports are disabled for this check")


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--date", type=dt.date.fromisoformat, default=dt.date.today())
    args = parser.parse_args()
    sys.meta_path.insert(0, BlockAccess())
    import psycopg
    from openpyxl import load_workbook
    from inspection_records_search.config import get_database_backend
    from inspection_records_search.app import _build_bridge
    from inspection_records_search.services.export_service import export_to_xlsx
    if get_database_backend() != "postgres":
        print("NOT RUN: DB_BACKEND must be postgres")
        return 1
    connect = psycopg.connect

    def read_only_connect(dsn):
        return connect(dsn, connect_timeout=5,
                       options="-c default_transaction_read_only=on -c statement_timeout=10000")

    try:
        with patch.object(psycopg, "connect", read_only_connect):
            bridge = _build_bridge()
            if bridge._bootstrap.startup_error or bridge._svc is None:
                raise RuntimeError("Startup configuration failed")
            svc = bridge._svc
            _, inspectors = svc.fetch_inspectors()
            svc.fetch_koutei_distinct_values()
            headers, rows = svc.fetch_main_detail(args.date, args.date, None)
            named = next((r for r in inspectors if r[0] and r[1]), None)
            if named is None:
                raise RuntimeError("No named inspector available for the inquiry check")
            svc.fetch_personal_records(str(named[0]).strip(), args.date, args.date)
            svc.fetch_personal_summary(str(named[0]).strip(), args.date, args.date)
            part = rows[0][headers.index("品番")] if rows else "__BUILD_SMOKE_NO_MATCH__"
            svc.fetch_lot_aggregate(part, None)
            with tempfile.TemporaryDirectory() as directory:
                output = export_to_xlsx(None, "smoke.xlsx", headers, rows, output_dir=Path(directory))
                book = load_workbook(output, read_only=True)
                try:
                    assert list(next(book.active.values)) == headers
                    assert book.active.max_row == len(rows) + 1
                finally:
                    book.close()
        print(f"PASS: PostgreSQL masters, detail, personal inquiry, lot aggregate, Excel; date={args.date}; detail_rows={len(rows)}")
        return 0
    except Exception as exc:
        # Do not print connection values or database records.
        print(f"FAIL: {type(exc).__name__}; connection settings and data are not logged")
        return 1


if __name__ == "__main__":
    raise SystemExit(main())
